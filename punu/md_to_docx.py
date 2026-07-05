"""Markdown → DOCX (schwarze Schrift). Usage: python md_to_docx.py input.md [output.docx]"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

BLACK = RGBColor(0, 0, 0)


def add_run(paragraph, text: str, *, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    run.font.color.rgb = BLACK
    run.font.size = Pt(11)
    run.bold = bold


def convert(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    text = md_path.read_text(encoding="utf-8")
    in_table = False
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows, in_table
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=cols)
        table.style = "Table Grid"
        for ri, row in enumerate(table_rows):
            for ci, cell in enumerate(row):
                if ci < cols:
                    p = table.rows[ri].cells[ci].paragraphs[0]
                    add_run(p, cell.strip())
        table_rows = []
        in_table = False

    for raw in text.splitlines():
        line = raw.rstrip()

        if line.startswith("|") and "|" in line[1:]:
            if re.match(r"^\|[\s\-:|]+\|$", line):
                continue
            cells = [c.strip() for c in line.strip("|").split("|")]
            table_rows.append(cells)
            in_table = True
            continue
        if in_table:
            flush_table()

        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            for run in p.runs:
                run.font.color.rgb = BLACK
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            for run in p.runs:
                run.font.color.rgb = BLACK
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            for run in p.runs:
                run.font.color.rgb = BLACK
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_run(p, line[2:].strip(), bold=True)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_run(p, line[2:].strip())
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_run(p, re.sub(r"^\d+\.\s", "", line).strip())
            continue

        p = doc.add_paragraph()
        parts = re.split(r"(\*\*.*?\*\*)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                add_run(p, part[2:-2], bold=True)
            else:
                add_run(p, part)

    flush_table()
    doc.save(str(docx_path))
    print(f"OK: {docx_path}")


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python md_to_docx.py input.md [output.docx]")
        sys.exit(1)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    convert(src, dst)


if __name__ == "__main__":
    main()
